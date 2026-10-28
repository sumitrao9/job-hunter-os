import os
import uuid
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from app.config import settings


def _ensure_upload_dir() -> str:
    os.makedirs(settings.upload_dir, exist_ok=True)
    return settings.upload_dir


def generate_resume_docx(
    profile: Any,
    job: Any,
    match_result: dict[str, Any] | None = None,
) -> str:
    """Generate a tailored resume DOCX and return the stored file path."""
    upload_dir = _ensure_upload_dir()
    filename = f"resume_{uuid.uuid4().hex}.docx"
    path = os.path.join(upload_dir, filename)

    doc = Document()

    # Header — name + contact
    name = (
        (profile.full_name or "Candidate")
        if hasattr(profile, "full_name")
        else "Candidate"
    )
    heading = doc.add_heading(name, level=1)
    heading.alignment = WD_ALIGN_PARAGRAPH.LEFT

    contact_parts = []
    if hasattr(profile, "email") and profile.email:
        contact_parts.append(str(profile.email))
    if hasattr(profile, "phone") and profile.phone:
        contact_parts.append(str(profile.phone))
    if hasattr(profile, "location") and profile.location:
        contact_parts.append(str(profile.location))
    if contact_parts:
        doc.add_paragraph(" | ".join(contact_parts))

    # Summary
    summary_text = None
    if hasattr(profile, "parsed_resume") and profile.parsed_resume:
        summary_text = profile.parsed_resume.get("summary")
    if not summary_text and match_result:
        summary_text = (
            f"Profile matched {match_result.get('score', 0)}% against "
            f"{job.title if hasattr(job, 'title') else 'target role'}."
        )
    if summary_text:
        doc.add_heading("Summary", level=2)
        doc.add_paragraph(str(summary_text))

    # Skills
    skills = []
    if hasattr(profile, "skills") and profile.skills:
        skills = profile.skills if isinstance(profile.skills, list) else []
    if match_result and match_result.get("matched_skills"):
        matched = set(match_result["matched_skills"]) - set(skills)
        if matched:
            skills = list(skills) + list(matched)[:10]
    if skills:
        doc.add_heading("Skills", level=2)
        for skill in skills[:30]:
            doc.add_paragraph(str(skill), style="List Bullet")

    # Experience
    experience = []
    if hasattr(profile, "experience") and profile.experience:
        experience = profile.experience if isinstance(profile.experience, list) else []
    if experience:
        doc.add_heading("Experience", level=2)
        for role in experience[:10]:
            if isinstance(role, dict):
                title = role.get("title", "")
                company = role.get("company", "")
                period = role.get("period", "")
                bullets = role.get("bullets", [])
            else:
                title = getattr(role, "title", "")
                company = getattr(role, "company", "")
                period = getattr(role, "period", "")
                bullets = getattr(role, "bullets", [])
            p = doc.add_paragraph()
            p.add_run(f"{title} at {company}").bold = True
            if period:
                p.add_run(f" ({period})")
            for bullet in bullets[:6]:
                doc.add_paragraph(str(bullet), style="List Bullet")

    # Education
    education = []
    if hasattr(profile, "education") and profile.education:
        education = profile.education if isinstance(profile.education, list) else []
    if education:
        doc.add_heading("Education", level=2)
        for edu in education[:5]:
            if isinstance(edu, dict):
                degree = edu.get("degree", "")
                school = edu.get("school", "")
                period = edu.get("period", "")
            else:
                degree = getattr(edu, "degree", "")
                school = getattr(edu, "school", "")
                period = getattr(edu, "period", "")
            text = f"{degree} from {school}" + (f" ({period})" if period else "")
            doc.add_paragraph(text, style="List Bullet")

    doc.save(path)
    return path


def generate_cover_letter_docx(
    profile: Any,
    job: Any,
    match_result: dict[str, Any] | None = None,
    body_text: str | None = None,
) -> str:
    """Generate a cover letter DOCX and return the stored file path."""
    upload_dir = _ensure_upload_dir()
    filename = f"cover_letter_{uuid.uuid4().hex}.docx"
    path = os.path.join(upload_dir, filename)

    doc = Document()
    name = (
        (profile.full_name or "Candidate")
        if hasattr(profile, "full_name")
        else "Candidate"
    )
    doc.add_heading(name, level=1)

    if body_text:
        doc.add_paragraph(str(body_text))
    else:
        job_title = getattr(job, "title", "the position")
        company = getattr(job, "company", "your team")
        intro = (
            f"Dear Hiring Manager,\n\n"
            f"I am writing to express my interest in {job_title} at {company}."
        )
        doc.add_paragraph(intro)
        doc.add_paragraph(
            "Based on my background and the requirements outlined, I believe I am well-suited "
            "to contribute to your team. I have attached my resume for your review and welcome "
            "the opportunity to discuss how I can add value."
        )
        doc.add_paragraph(
            "Thank you for considering my application. I look forward to hearing from you."
        )
        doc.add_paragraph("Sincerely,\n" + name)

    doc.save(path)
    return path


def docx_to_bytes(path: str) -> bytes:
    with open(path, "rb") as f:
        return f.read()
